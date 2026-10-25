"""Multi-sheet Excel (.xlsx) workbook exporter and reader for MALDI-GlycoID (Stage 10).

Uses `openpyxl` when installed, with a standards-compliant ECMA-376 OpenXML `.xlsx`
generator and reader fallback so multi-sheet `.xlsx` exports work 100% offline in
any Python 3.11+ environment and open cleanly in Microsoft Excel, LibreOffice Calc,
and pandas/openpyxl.
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Mapping, Sequence
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape as xml_escape
import zipfile


def _sanitize_sheet_name(name: str, existing: set[str]) -> str:
    """Sanitize Excel worksheet title (max 31 chars, no `: \\ / ? * [ ]`)."""
    cleaned = re.sub(r"[:\\/?*\[\]]+", "_", name).strip() or "Sheet"
    base = cleaned[:28]
    candidate = cleaned[:31]
    counter = 1
    while candidate in existing:
        counter += 1
        candidate = f"{base}_{counter}"[:31]
    existing.add(candidate)
    return candidate


def _col_index_to_excel_letters(col_idx_1based: int) -> str:
    """Convert 1-based column index to Excel column letters (1 -> A, 27 -> AA)."""
    letters: List[str] = []
    n = col_idx_1based
    while n > 0:
        n, rem = divmod(n - 1, 26)
        letters.append(chr(65 + rem))
    return "".join(reversed(letters))


def _format_cell_value(val: Any) -> Any:
    if val is None:
        return ""
    if isinstance(val, bool):
        return "TRUE" if val else "FALSE"
    if isinstance(val, (int, float)):
        if val != val or val in (float("inf"), float("-inf")):  # NaN or Inf
            return ""
        return val
    if isinstance(val, (list, tuple)):
        return "; ".join(str(x) for x in val)
    if isinstance(val, dict):
        return "; ".join(f"{k}={v}" for k, v in val.items())
    return str(val)


def _build_worksheet_xml(rows: Sequence[Mapping[str, Any]]) -> str:
    """Build ECMA-376 `xl/worksheets/sheetN.xml` with header row and typed cells."""
    headers: List[str] = []
    seen: set[str] = set()
    for r in rows:
        for k in r.keys():
            if k not in seen:
                seen.add(k)
                headers.append(k)

    xml_lines: List[str] = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">',
        "  <sheetData>",
    ]

    if headers:
        # Row 1: Header row (style s="1" bold)
        xml_lines.append('    <row r="1">')
        for c_idx, h in enumerate(headers, start=1):
            col_ref = f"{_col_index_to_excel_letters(c_idx)}1"
            safe_h = xml_escape(str(h))
            xml_lines.append(
                f'      <c r="{col_ref}" s="1" t="inlineStr"><is><t>{safe_h}</t></is></c>'
            )
        xml_lines.append("    </row>")

        # Data rows
        for r_idx, row_dict in enumerate(rows, start=2):
            xml_lines.append(f'    <row r="{r_idx}">')
            for c_idx, h in enumerate(headers, start=1):
                col_ref = f"{_col_index_to_excel_letters(c_idx)}{r_idx}"
                cell_val = _format_cell_value(row_dict.get(h))
                if cell_val == "":
                    continue
                if isinstance(cell_val, (int, float)) and not isinstance(cell_val, bool):
                    xml_lines.append(f'      <c r="{col_ref}" t="n"><v>{cell_val}</v></c>')
                else:
                    safe_txt = xml_escape(str(cell_val))
                    xml_lines.append(
                        f'      <c r="{col_ref}" t="inlineStr"><is><t>{safe_txt}</t></is></c>'
                    )
            xml_lines.append("    </row>")

    xml_lines.append("  </sheetData>")
    xml_lines.append("</worksheet>")
    return "\n".join(xml_lines)


def export_tables_to_xlsx_bytes(
    sheets: Mapping[str, Sequence[Mapping[str, Any]]],
) -> bytes:
    """Export a dictionary of `{sheet_name: rows}` into a multi-sheet `.xlsx` workbook."""
    if not sheets:
        sheets = {"Summary": [{"info": "No data rows provided"}]}

    try:
        import openpyxl  # type: ignore

        wb = openpyxl.Workbook()
        # Remove default sheet after creating named sheets
        default_ws = wb.active
        used_names: set[str] = set()
        for raw_name, rows in sheets.items():
            safe_title = _sanitize_sheet_name(raw_name, used_names)
            ws = wb.create_sheet(title=safe_title)
            headers: List[str] = []
            seen: set[str] = set()
            for r in rows:
                for k in r.keys():
                    if k not in seen:
                        seen.add(k)
                        headers.append(k)
            if headers:
                ws.append(headers)
                for r in rows:
                    ws.append([_format_cell_value(r.get(h)) for h in headers])
        if default_ws is not None and len(wb.sheetnames) > 1:
            wb.remove(default_ws)
        out_buf = io.BytesIO()
        wb.save(out_buf)
        return out_buf.getvalue()
    except ImportError:
        pass

    # Standards-compliant ECMA-376 OpenXML .xlsx writer
    buf = io.BytesIO()
    used_names = set()
    sheet_entries: List[ tuple[int, str, Sequence[Mapping[str, Any]]] ] = []
    for idx, (raw_name, rows) in enumerate(sheets.items(), start=1):
        safe_title = _sanitize_sheet_name(raw_name, used_names)
        sheet_entries.append((idx, safe_title, rows))

    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        # 1. [Content_Types].xml
        ct_lines = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
            '  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
            '  <Default Extension="xml" ContentType="application/xml"/>',
            '  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>',
            '  <Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>',
        ]
        for idx, _, _ in sheet_entries:
            ct_lines.append(
                f'  <Override PartName="/xl/worksheets/sheet{idx}.xml" '
                f'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            )
        ct_lines.append("</Types>")
        zf.writestr("[Content_Types].xml", "\n".join(ct_lines))

        # 2. _rels/.rels
        rels_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n'
            '  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>\n'
            "</Relationships>"
        )
        zf.writestr("_rels/.rels", rels_xml)

        # 3. xl/workbook.xml
        wb_lines = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">',
            "  <sheets>",
        ]
        for idx, title, _ in sheet_entries:
            wb_lines.append(
                f'    <sheet name="{xml_escape(title)}" sheetId="{idx}" r:id="rId{idx}"/>'
            )
        wb_lines.append("  </sheets>")
        wb_lines.append("</workbook>")
        zf.writestr("xl/workbook.xml", "\n".join(wb_lines))

        # 4. xl/_rels/workbook.xml.rels
        wb_rels_lines = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">',
        ]
        for idx, _, _ in sheet_entries:
            wb_rels_lines.append(
                f'  <Relationship Id="rId{idx}" '
                f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
                f'Target="worksheets/sheet{idx}.xml"/>'
            )
        styles_rid = len(sheet_entries) + 1
        wb_rels_lines.append(
            f'  <Relationship Id="rId{styles_rid}" '
            f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
            f'Target="styles.xml"/>'
        )
        wb_rels_lines.append("</Relationships>")
        zf.writestr("xl/_rels/workbook.xml.rels", "\n".join(wb_rels_lines))

        # 5. xl/styles.xml (minimal valid stylesheet with normal font 0 and bold font 1)
        styles_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">\n'
            '  <fonts count="2">\n'
            '    <font><sz val="10"/><name val="Calibri"/></font>\n'
            '    <font><b/><sz val="10"/><name val="Calibri"/></font>\n'
            "  </fonts>\n"
            '  <fills count="1"><fill><patternFill patternType="none"/></fill></fills>\n'
            '  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>\n'
            '  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>\n'
            '  <cellXfs count="2">\n'
            '    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>\n'
            '    <xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/>\n'
            "  </cellXfs>\n"
            "</styleSheet>"
        )
        zf.writestr("xl/styles.xml", styles_xml)

        # 6. xl/worksheets/sheet{idx}.xml
        for idx, _, rows in sheet_entries:
            zf.writestr(f"xl/worksheets/sheet{idx}.xml", _build_worksheet_xml(rows))

    return buf.getvalue()


def _excel_letters_to_col_index(col_letters: str) -> int:
    """Convert Excel column letters (e.g. 'A' -> 0, 'Z' -> 25, 'AA' -> 26)."""
    idx = 0
    for ch in col_letters:
        if "A" <= ch <= "Z":
            idx = idx * 26 + (ord(ch) - 64)
    return idx - 1


def read_xlsx_bytes_sheets(xlsx_bytes: bytes) -> Dict[str, List[Dict[str, Any]]]:
    """Parse a multi-sheet `.xlsx` byte stream back into `{sheet_name: [row_dicts]}`."""
    ns = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    result: Dict[str, List[Dict[str, Any]]] = {}

    with zipfile.ZipFile(io.BytesIO(xlsx_bytes), "r") as zf:
        wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
        rels_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))

        rid_to_target: Dict[str, str] = {}
        for rel in rels_root.findall("pkgrel:Relationship", ns):
            r_id = rel.attrib.get("Id", "")
            target = rel.attrib.get("Target", "").lstrip("/")
            if not target.startswith("xl/"):
                target = f"xl/{target}"
            rid_to_target[r_id] = target

        # Optional sharedStrings table
        shared_strings: List[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            ss_root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in ss_root.findall("main:si", ns):
                texts = [t.text or "" for t in si.findall(".//main:t", ns)]
                shared_strings.append("".join(texts))

        for sheet_el in wb_root.findall(".//main:sheet", ns):
            sheet_name = sheet_el.attrib.get("name", "Sheet")
            r_id = sheet_el.attrib.get(f"{{{ns['r']}}}id", "")
            target_path = rid_to_target.get(r_id)
            if not target_path or target_path not in zf.namelist():
                continue

            ws_root = ET.fromstring(zf.read(target_path))
            parsed_rows: List[Dict[int, Any]] = []
            max_col = -1

            for row_el in ws_root.findall(".//main:sheetData/main:row", ns):
                row_map: Dict[int, Any] = {}
                for c_el in row_el.findall("main:c", ns):
                    ref = c_el.attrib.get("r", "A1")
                    col_letters = "".join(ch for ch in ref if ch.isalpha())
                    c_idx = _excel_letters_to_col_index(col_letters)
                    if c_idx > max_col:
                        max_col = c_idx

                    c_type = c_el.attrib.get("t", "n")
                    val: Any = None
                    if c_type == "inlineStr":
                        t_el = c_el.find(".//main:t", ns)
                        val = t_el.text if t_el is not None and t_el.text is not None else ""
                    elif c_type == "s":
                        v_el = c_el.find("main:v", ns)
                        if v_el is not None and v_el.text is not None:
                            s_idx = int(v_el.text)
                            val = shared_strings[s_idx] if 0 <= s_idx < len(shared_strings) else ""
                    else:
                        v_el = c_el.find("main:v", ns)
                        if v_el is not None and v_el.text is not None:
                            num_txt = v_el.text
                            try:
                                val = int(num_txt) if "." not in num_txt and "e" not in num_txt.lower() else float(num_txt)
                            except ValueError:
                                val = num_txt
                    row_map[c_idx] = val
                parsed_rows.append(row_map)

            if not parsed_rows or max_col < 0:
                result[sheet_name] = []
                continue

            header_map = parsed_rows[0]
            headers = [str(header_map.get(c, f"Col_{c+1}")) for c in range(max_col + 1)]
            sheet_dicts: List[Dict[str, Any]] = []
            for r_map in parsed_rows[1:]:
                sheet_dicts.append({headers[c]: r_map.get(c) for c in range(max_col + 1)})
            result[sheet_name] = sheet_dicts

    return result
